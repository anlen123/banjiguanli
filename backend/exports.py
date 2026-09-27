"""Self-contained Excel / portable ZIP exports, with no authenticated web URLs."""

from __future__ import annotations

import hashlib
import io
import json
import re
import warnings
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import cast
from urllib.parse import quote
from xml.etree import ElementTree as ET

from ole_package import package_file
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Color, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image, ImageDraw, ImageOps, UnidentifiedImageError

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
VML = "urn:schemas-microsoft-com:vml"
OFFICE = "urn:schemas-microsoft-com:office:office"
EXCEL = "urn:schemas-microsoft-com:office:excel"
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp"}
MAX_EXPORT_BYTES = 100 * 1024 * 1024
MAX_EXPORT_FILES = 500


@dataclass
class Asset:
    id: str
    user_id: int
    name: str
    question: str
    filename: str
    data: bytes
    path: str


def portable_name(name: str) -> str:
    value = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "_", name).strip(" .")
    value = value.encode("utf-8")[:120].decode("utf-8", errors="ignore")
    return value or "attachment"


def text(value) -> str:
    return ILLEGAL_CHARACTERS_RE.sub("", str(value if value is not None else ""))


def thumbnail(asset: Asset):
    if Path(asset.filename).suffix.lower() not in IMAGE_SUFFIXES:
        return None, "非图片，请打开原文件"
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(io.BytesIO(asset.data)) as source:
                if source.width * source.height > 20_000_000:
                    return None, "图片过大不生成预览，原文件已保留"
                source.seek(0)
                image = ImageOps.exif_transpose(source).convert("RGB")
                image.thumbnail((220, 150))
                output = io.BytesIO()
                image.save(output, "PNG")
                output.seek(0)
                return output, "图片预览（原文件另附，动图预览首帧）"
    except (
        UnidentifiedImageError,
        OSError,
        ValueError,
        Image.DecompressionBombError,
        Image.DecompressionBombWarning,
    ):
        return None, "无法生成图片预览，原文件已保留"


def literal_rows(sheet):
    for row in sheet:
        for cell in row:
            if cell.value is not None:
                cell.value = text(cell.value)
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in next(sheet.iter_rows()):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=Color(rgb="427B69"))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def build_workbook(
    project: dict, records: list[dict], assets: list[Asset], embedded: bool
) -> bytes:
    book = Workbook()
    results = cast(Worksheet, book.active)
    results.title = "填写结果"
    files = book.create_sheet("附件")
    instructions = book.create_sheet("使用说明")
    results.append(
        ["姓名", "学号", "更新时间 (UTC)"] + [text(q["label"]) for q in project["fields"]]
    )
    by_id = {a.id: (i + 2, a) for i, a in enumerate(assets)}
    for record in records:
        answers = record["answers"]
        row = [text(record["name"]), text(record["student_no"]), record["updated_at"]]
        links = []
        for index, question in enumerate(project["fields"], 4):
            value = answers.get(question["id"], "")
            if isinstance(value, list):
                row.append(
                    "\n".join(
                        f"附件 {by_id[aid][0] - 1}: {text(by_id[aid][1].filename)}" for aid in value
                    )
                )
                if value:
                    links.append((index, by_id[value[0]][0]))
            else:
                row.append(text(value))
        results.append(row)
        for column, file_row in links:
            cell = results.cell(results.max_row, column)
            cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'附件'!A{file_row}")
        results.row_dimensions[results.max_row].height = min(
            120,
            max(32, 18 * max((len(v) for v in answers.values() if isinstance(v, list)), default=1)),
        )
    for i in range(1, results.max_column + 1):
        results.column_dimensions[get_column_letter(i)].width = 28
    files.append(
        ["编号", "填写人", "收集问题", "原文件名", "大小 (KB)", "图片预览", "原始附件 / 打开方式"]
    )
    for i, asset in enumerate(assets, 2):
        preview, message = thumbnail(asset)
        files.append(
            [
                str(i - 1),
                text(asset.name),
                text(asset.question),
                text(asset.filename),
                f"{len(asset.data) / 1024:.1f}",
                message if preview is None else "",
                "双击下方附件图标" if embedded else "点击打开原文件",
            ]
        )
        if preview:
            files.add_image(ExcelImage(preview), f"F{i}")
        if not embedded:
            files.cell(i, 7).hyperlink = Hyperlink(ref=f"G{i}", target=quote(asset.path, safe="/"))
        files.row_dimensions[i].height = 140
    for column, width in zip("ABCDEFG", (8, 18, 22, 40, 14, 34, 22), strict=True):
        files.column_dimensions[column].width = width
    files.row_dimensions[1].height = 26
    instructions.append(["说明", "内容"])
    instructions.append(["收集项目", text(project["title"])])
    instructions.append(
        [
            "离线文件",
            "本文件不依赖系统登录或服务器下载地址。附件表包含全部已提交附件；未提交的上传不导出。",
        ]
    )
    instructions.append(
        [
            "图片",
            "附件表直接展示缩略图；完整原文件保持不变。动图只预览首帧，异常或过大图片仅保留原文件。",
        ]
    )
    instructions.append(
        [
            "打开附件",
            "Excel 嵌入对象需使用支持 OLE 的 Windows 桌面 Excel，双击图标打开；需安装对应文件阅读程序。Office 安全策略可能阻止打开，请勿为此关闭安全保护。"
            if embedded
            else "请先完整解压 ZIP，再打开 workbook.xlsx；原文件在 attachments 文件夹，保留相对目录即可离线访问。",
        ]
    )
    instructions.append(
        [
            "兼容说明",
            "网页版 Excel、Mac Excel、WPS、手机表格对 OLE 支持不同；打不开嵌入对象时，请改用系统的 ZIP 导出。不要使用不支持 OLE 的软件另存为覆盖唯一备份。",
        ]
    )
    instructions.append(
        [
            "文件命名",
            "OLE 对象内部使用编号英文文件名以兼容系统编码，原文件名在附件表；ZIP 文件名包含唯一 ID，防止同名覆盖。",
        ]
    )
    instructions.append(
        ["数据保护", "导出包含实际个人资料和文件，无需登录即可阅读；请妥善保管，不要公开分享。"]
    )
    instructions.column_dimensions["A"].width = 20
    instructions.column_dimensions["B"].width = 100
    for sheet in book.worksheets:
        literal_rows(sheet)
    results.freeze_panes = "D2"
    output = io.BytesIO()
    book.save(output)
    return embed_objects(output.getvalue(), assets) if embedded and assets else output.getvalue()


def icon_png() -> bytes:
    image = Image.new("RGB", (80, 72), "#f5f8ed")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((18, 5, 61, 53), radius=5, fill="white", outline="#43836d", width=2)
    for y in (19, 29, 39):
        draw.line((27, y, 53, y), fill="#8cab7c", width=2)
    draw.text((25, 56), "OPEN", fill="#43836d")
    out = io.BytesIO()
    image.save(out, "PNG")
    return out.getvalue()


def xml_bytes(root):
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def embed_objects(workbook: bytes, assets: list[Asset]) -> bytes:
    """Attach Packager storage, OLE relationships and visible VML icon shapes."""
    with zipfile.ZipFile(io.BytesIO(workbook)) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    sheet_path = "xl/worksheets/sheet2.xml"
    rels_path = "xl/worksheets/_rels/sheet2.xml.rels"
    sheet = ET.fromstring(parts[sheet_path])
    rels = (
        ET.fromstring(parts[rels_path])
        if rels_path in parts
        else ET.Element(f"{{{PKG}}}Relationships")
    )
    ET.SubElement(
        rels,
        f"{{{PKG}}}Relationship",
        Id="rIdAttachmentVML",
        Type=f"{REL}/vmlDrawing",
        Target="../drawings/attachments.vml",
    )
    ET.SubElement(sheet, f"{{{MAIN}}}legacyDrawing", {f"{{{REL}}}id": "rIdAttachmentVML"})
    objects = ET.SubElement(sheet, f"{{{MAIN}}}oleObjects")
    vml = ET.Element("xml")
    layout = ET.SubElement(vml, f"{{{OFFICE}}}shapelayout", {f"{{{VML}}}ext": "edit"})
    ET.SubElement(layout, f"{{{OFFICE}}}idmap", {f"{{{VML}}}ext": "edit", "data": "1"})
    shape_type = ET.SubElement(
        vml,
        f"{{{VML}}}shapetype",
        {
            "id": "_x0000_t75",
            "coordsize": "21600,21600",
            f"{{{OFFICE}}}spt": "75",
            f"{{{OFFICE}}}preferrelative": "t",
            "path": "m@4@5l@4@11@9@11@9@5xe",
            "filled": "f",
            "stroked": "f",
        },
    )
    ET.SubElement(shape_type, f"{{{VML}}}stroke", joinstyle="miter")
    formulas = ET.SubElement(shape_type, f"{{{VML}}}formulas")
    for equation in [
        "if lineDrawn pixelLineWidth 0",
        "sum @0 1 0",
        "sum 0 0 @1",
        "prod @2 1 2",
        "prod @3 21600 pixelWidth",
        "prod @3 21600 pixelHeight",
        "sum @0 0 1",
        "prod @6 1 2",
        "prod @7 21600 pixelWidth",
        "sum @8 21600 0",
        "prod @7 21600 pixelHeight",
        "sum @10 21600 0",
    ]:
        ET.SubElement(formulas, f"{{{VML}}}f", eqn=equation)
    ET.SubElement(
        shape_type,
        f"{{{VML}}}path",
        {
            f"{{{OFFICE}}}extrusionok": "f",
            "gradientshapeok": "t",
            f"{{{OFFICE}}}connecttype": "rect",
        },
    )
    ET.SubElement(shape_type, f"{{{OFFICE}}}lock", {f"{{{VML}}}ext": "edit", "aspectratio": "t"})
    for i, asset in enumerate(assets, 1):
        rid, shape = f"rIdEmbedded{i}", str(1024 + i)
        name = f"attachment-{i}{Path(asset.filename).suffix.lower()}"
        parts[f"xl/embeddings/attachment{i}.bin"] = package_file(name, asset.data)
        ET.SubElement(
            rels,
            f"{{{PKG}}}Relationship",
            Id=rid,
            Type=f"{REL}/oleObject",
            Target=f"../embeddings/attachment{i}.bin",
        )
        ET.SubElement(
            objects,
            f"{{{MAIN}}}oleObject",
            {
                "progId": "Package",
                "dvAspect": "DVASPECT_ICON",
                "autoLoad": "0",
                "shapeId": shape,
                f"{{{REL}}}id": rid,
            },
        )
        node = ET.SubElement(
            vml,
            f"{{{VML}}}shape",
            {
                "id": f"_x0000_s{shape}",
                "type": "#_x0000_t75",
                "style": f"position:absolute;margin-left:710pt;margin-top:{40 + (i - 1) * 140}pt;width:60pt;height:54pt;z-index:{i}",
                f"{{{OFFICE}}}ole": "",
                "filled": "f",
                "stroked": "f",
            },
        )
        ET.SubElement(
            node,
            f"{{{VML}}}imagedata",
            {f"{{{OFFICE}}}relid": "rIdIcon", f"{{{OFFICE}}}title": text(asset.filename)},
        )
        client = ET.SubElement(node, f"{{{EXCEL}}}ClientData", ObjectType="Pict")
        ET.SubElement(client, f"{{{EXCEL}}}SizeWithCells")
        ET.SubElement(client, f"{{{EXCEL}}}Anchor").text = f"6, 12, {i}, 35, 6, 92, {i}, 107"
        ET.SubElement(client, f"{{{EXCEL}}}CF").text = "Pict"
        ET.SubElement(client, f"{{{EXCEL}}}AutoPict")
    vml_rels = ET.Element(f"{{{PKG}}}Relationships")
    ET.SubElement(
        vml_rels,
        f"{{{PKG}}}Relationship",
        Id="rIdIcon",
        Type=f"{REL}/image",
        Target="../media/attachment-icon.png",
    )
    parts["xl/drawings/attachments.vml"] = xml_bytes(vml)
    parts["xl/drawings/_rels/attachments.vml.rels"] = xml_bytes(vml_rels)
    parts["xl/media/attachment-icon.png"] = icon_png()
    content_types = ET.fromstring(parts["[Content_Types].xml"])
    for extension, mime in [
        ("bin", "application/vnd.openxmlformats-officedocument.oleObject"),
        ("vml", "application/vnd.openxmlformats-officedocument.vmlDrawing"),
        ("png", "image/png"),
    ]:
        if not any(n.get("Extension") == extension for n in content_types):
            ET.SubElement(content_types, f"{{{CT}}}Default", Extension=extension, ContentType=mime)
    parts["[Content_Types].xml"] = xml_bytes(content_types)
    parts[sheet_path], parts[rels_path] = xml_bytes(sheet), xml_bytes(rels)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return out.getvalue()


def build_zip(project: dict, records: list[dict], assets: list[Asset]) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("workbook.xlsx", build_workbook(project, records, assets, embedded=False))
        manifest = []
        for asset in assets:
            archive.writestr(asset.path, asset.data)
            manifest.append(
                {
                    "id": asset.id,
                    "user_id": asset.user_id,
                    "name": asset.name,
                    "question": asset.question,
                    "filename": asset.filename,
                    "path": asset.path,
                    "bytes": len(asset.data),
                    "sha256": hashlib.sha256(asset.data).hexdigest(),
                }
            )
        archive.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))
        archive.writestr(
            "README.txt",
            "请完整解压 ZIP 后打开 workbook.xlsx；图片在附件表预览，原文件在 attachments 目录。\n保持目录结构可使用 Excel 中的相对链接；软件拦截链接时可直接打开目录中的文件。\nmanifest.json 记录原文件名、填写人、相对路径及 SHA-256。\n本压缩包包含实际学生资料，请妥善保管，勿公开分享。\n",
        )
    return output.getvalue()
