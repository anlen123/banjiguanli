import io
from datetime import datetime, timedelta, timezone

import app as service
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

HEADERS = {"X-Requested-With": "class-island"}


@pytest.fixture()
def clients(tmp_path, monkeypatch):
    monkeypatch.setattr(service, "DATA", tmp_path)
    service.create_admin("admin", "班长", "administrator123")
    with (
        TestClient(service.app, headers=HEADERS) as root,
        TestClient(service.app, headers=HEADERS) as student,
        TestClient(service.app, headers=HEADERS) as other,
    ):
        assert (
            root.post(
                "/api/auth/login",
                json={"username": "admin", "password": "administrator123"},
            ).status_code
            == 200
        )
        for client, username in [(student, "student"), (other, "other")]:
            assert (
                client.post(
                    "/api/auth/register",
                    json={
                        "username": username,
                        "password": "student123",
                        "name": username,
                        "student_no": "001",
                    },
                ).status_code
                == 201
            )
            assert (
                client.post(
                    "/api/auth/login",
                    json={"username": username, "password": "student123"},
                ).status_code
                == 403
            )
            uid = next(u["id"] for u in root.get("/api/users").json() if u["username"] == username)
            assert (
                root.patch(f"/api/users/{uid}/status", json={"status": "approved"}).status_code
                == 200
            )
            assert (
                client.post(
                    "/api/auth/login",
                    json={"username": username, "password": "student123"},
                ).status_code
                == 200
            )
        yield root, student, other


def make_project(root, **kwargs):
    body = {
        "title": "奖学金资料收集",
        "description": "请提交资料",
        "published": True,
        "fields": [
            {"id": "name", "label": "姓名", "required": True, "type": "text"},
            {"id": "file", "label": "证明", "type": "file"},
        ],
    }
    body.update(kwargs)
    response = root.post("/api/projects", json=body)
    assert response.status_code == 201, response.text
    return response.json(), body


def test_approval_and_permissions(clients):
    root, student, _ = clients
    assert student.get("/api/users").status_code == 403
    assert student.get("/api/ai/config").status_code == 403
    assert student.put("/api/ai/config", json={}).status_code == 403
    assert student.post("/api/projects", json={}).status_code == 403
    uid = student.get("/api/auth/me").json()["id"]
    root.patch(f"/api/users/{uid}/status", json={"status": "rejected"})
    assert student.get("/api/auth/me").status_code == 401
    assert root.patch("/api/users/1/status", json={"status": "rejected"}).status_code == 400


def test_submission_privacy_files_and_export(clients):
    root, student, other = clients
    p, _ = make_project(root)
    pid = p["id"]
    uid = student.get("/api/auth/me").json()["id"]
    assert student.put(f"/api/projects/{pid}/submission", json={"answers": {}}).status_code == 422
    upload = student.post(
        f"/api/projects/{pid}/files/file",
        files={"file": ("成绩.pdf", b"%PDF-example", "application/pdf")},
    )
    assert upload.status_code == 201, upload.text
    fid = upload.json()["id"]
    assert (
        student.post(
            f"/api/projects/{pid}/files/file", files={"file": ("bad.html", b"<script>")}
        ).status_code
        == 422
    )
    assert other.get(f"/api/files/{fid}").status_code == 403
    assert root.get(f"/api/files/{fid}").content == b"%PDF-example"
    assert (
        student.put(
            f"/api/projects/{pid}/submission",
            json={"answers": {"name": "=SUM(1,2)", "file": [fid]}},
        ).status_code
        == 200
    )
    assert (
        other.put(
            f"/api/projects/{pid}/submission",
            json={"answers": {"name": "other", "file": [fid]}},
        ).status_code
        == 422
    )
    assert other.get(f"/api/projects/{pid}/submission?user_id={uid}").status_code == 403
    assert (
        other.put(
            f"/api/projects/{pid}/submission?user_id={uid}",
            json={"answers": {"name": "x"}},
        ).status_code
        == 403
    )
    assert student.get(f"/api/projects/{pid}/submissions").status_code == 403
    assert student.get(f"/api/projects/{pid}/export").status_code == 403
    result = root.get(f"/api/projects/{pid}/submissions").json()
    assert result[0]["files"][fid]["filename"] == "成绩.pdf"
    exported = root.get(f"/api/projects/{pid}/export")
    assert exported.status_code == 200
    book = load_workbook(io.BytesIO(exported.content))
    from openpyxl.worksheet.worksheet import Worksheet

    sheet = book["填写结果"]
    assert isinstance(sheet, Worksheet)
    assert sheet.cell(2, 4).value == "=SUM(1,2)"
    assert sheet.cell(2, 4).data_type == "s"
    assert "成绩.pdf" in str(sheet.cell(2, 5).value)


def test_deadline_admin_override_and_reopen(clients):
    root, student, _ = clients
    p, body = make_project(
        root, deadline=(datetime.now(timezone.utc) - timedelta(hours=1)).isoformat()
    )
    pid = p["id"]
    uid = student.get("/api/auth/me").json()["id"]
    answers = {"answers": {"name": "student"}}
    assert student.put(f"/api/projects/{pid}/submission", json=answers).status_code == 403
    assert (
        student.post(f"/api/projects/{pid}/files/file", files={"file": ("x.pdf", b"x")}).status_code
        == 403
    )
    assert (
        root.put(f"/api/projects/{pid}/submission?user_id={uid}", json=answers).status_code == 200
    )
    body["deadline"] = (datetime.now(timezone.utc) + timedelta(days=1)).isoformat()
    assert root.put(f"/api/projects/{pid}", json=body).status_code == 200
    assert student.put(f"/api/projects/{pid}/submission", json=answers).status_code == 200
    body["fields"] = []
    assert root.put(f"/api/projects/{pid}", json=body).status_code == 422
    body["fields"] = [{"id": "different", "label": "新问题"}]
    assert root.put(f"/api/projects/{pid}", json=body).status_code == 409


def test_drafts_and_password_revocation(clients):
    root, student, _ = clients
    p, _ = make_project(root, published=False)
    assert student.get("/api/projects").json() == []
    assert student.get(f"/api/projects/{p['id']}").status_code == 404
    uid = student.get("/api/auth/me").json()["id"]
    assert (
        root.put(f"/api/users/{uid}/password", json={"password": "changed123"}).status_code == 200
    )
    assert student.get("/api/auth/me").status_code == 401
    assert (
        student.post(
            "/api/auth/login", json={"username": "student", "password": "student123"}
        ).status_code
        == 401
    )
    assert (
        student.post(
            "/api/auth/login", json={"username": "student", "password": "changed123"}
        ).status_code
        == 200
    )
    assert (
        student.put(
            "/api/auth/password",
            json={"old_password": "changed123", "password": "newer12345"},
        ).status_code
        == 200
    )
    assert student.get("/api/auth/me").status_code == 401


def test_csrf_and_ai_validation(clients):
    root, student, _ = clients
    assert root.post("/api/auth/logout", headers={"X-Requested-With": ""}).status_code == 403
    assert (
        root.post("/api/auth/logout", headers={"Origin": "https://evil.invalid"}).status_code == 403
    )
    assert root.put("/api/ai/config", json={"base_url": "http://localhost/v1"}).status_code == 422
    assert root.put("/api/ai/config", json={"base_url": "https://127.0.0.1/v1"}).status_code == 422
    assert (
        root.post(
            "/api/ai/chat", json={"messages": [{"role": "user", "content": "你好"}]}
        ).status_code
        == 400
    )
    assert (
        student.post(
            "/api/ai/chat", json={"messages": [{"role": "user", "content": "你好"}]}
        ).status_code
        == 403
    )


def test_question_validation_and_file_boundaries(clients, monkeypatch):
    root, student, _ = clients
    p, _ = make_project(
        root,
        fields=[
            {"id": "num", "label": "数字", "type": "number", "required": True},
            {"id": "date", "label": "日期", "type": "date"},
            {"id": "choice", "label": "选择", "type": "select", "options": ["A", "B"]},
            {"id": "file", "label": "附件", "type": "file"},
        ],
    )
    pid = p["id"]
    for answers in [
        {"num": "NaN"},
        {"num": "1", "date": "not-date"},
        {"num": "1", "choice": "C"},
        {"num": "1", "bad": "x"},
        {"num": " "},
    ]:
        assert (
            student.put(f"/api/projects/{pid}/submission", json={"answers": answers}).status_code
            == 422
        )
    assert (
        student.put(
            f"/api/projects/{pid}/submission",
            json={"answers": {"num": "0", "date": "2026-01-01", "choice": "A"}},
        ).status_code
        == 200
    )
    monkeypatch.setattr(service, "MAX_FILE", 2)
    assert (
        student.post(
            f"/api/projects/{pid}/files/file", files={"file": ("a.pdf", b"123")}
        ).status_code
        == 413
    )
    assert list((service.DATA / "uploads").iterdir()) == []
    assert (
        student.post(
            f"/api/projects/{pid}/files/file", files={"file": ("empty.pdf", b"")}
        ).status_code
        == 422
    )


def test_ai_key_encryption_and_compatible_chat(clients, monkeypatch):
    root, _, _ = clients
    monkeypatch.setattr(service, "validate_ai_url", lambda url: "8.8.8.8")
    assert (
        root.put(
            "/api/ai/config",
            json={
                "base_url": "https://example.com/v1",
                "model": "test-model",
                "api_key": "super-secret-test",
            },
        ).status_code
        == 200
    )
    config = root.get("/api/ai/config").json()
    assert config["has_key"] is True and "super-secret-test" not in str(config)
    with service.db() as c:
        assert "super-secret-test" not in c.execute("SELECT value FROM settings").fetchone()[0]

    class MockClient:
        def __init__(self, **kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            pass

        async def post(self, url, headers, json, extensions):
            assert url == "https://8.8.8.8/v1/chat/completions"
            assert extensions == {"sni_hostname": "example.com"}
            assert headers["Host"] == "example.com"
            assert headers["Authorization"] == "Bearer super-secret-test"
            assert json["model"] == "test-model"
            assert "tools" not in json
            import httpx

            return httpx.Response(
                200,
                json={"choices": [{"message": {"content": "你好，班长"}}]},
                request=httpx.Request("POST", url),
            )

    monkeypatch.setattr(service.httpx, "AsyncClient", MockClient)
    assert (
        root.post("/api/ai/chat", json={"messages": [{"role": "user", "content": "你好"}]}).json()[
            "content"
        ]
        == "你好，班长"
    )
    assert root.put("/api/ai/config", json={"clear_key": True}).status_code == 200
    assert root.get("/api/ai/config").json()["has_key"] is False


def test_ai_destination_validation(monkeypatch):
    from fastapi import HTTPException

    monkeypatch.setattr(
        service.socket, "getaddrinfo", lambda *a, **k: [(2, 1, 6, "", ("8.8.8.8", 443))]
    )
    assert service.validate_ai_url("https://example.com/v1") == "8.8.8.8"
    for url in [
        "http://example.com",
        "https://example.com:bad",
        "https://[broken",
        "https://name:pass@example.com/v1",
        "https://example.com:8080/v1",
        "https://example.com/v1?key=x",
        "https://example.com/v1#fragment",
    ]:
        with pytest.raises(HTTPException) as error:
            service.validate_ai_url(url)
        assert error.value.status_code == 422
    monkeypatch.setattr(
        service.socket,
        "getaddrinfo",
        lambda *a, **k: [(2, 1, 6, "", ("8.8.8.8", 443)), (2, 1, 6, "", ("127.0.0.1", 443))],
    )
    with pytest.raises(HTTPException) as error:
        service.validate_ai_url("https://example.com/v1")
    assert error.value.status_code == 422


def test_export_control_characters(clients):
    root, student, _ = clients
    p, _ = make_project(root)
    assert (
        student.put(
            f"/api/projects/{p['id']}/submission", json={"answers": {"name": "hello\u0000world"}}
        ).status_code
        == 200
    )
    response = root.get(f"/api/projects/{p['id']}/export")
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    from openpyxl.worksheet.worksheet import Worksheet

    sheet = book.worksheets[0]
    assert isinstance(sheet, Worksheet)
    assert sheet.cell(2, 4).value == "helloworld"
