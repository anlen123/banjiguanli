import hashlib
import io
import json
import posixpath
import struct
import zipfile
from pathlib import PurePosixPath
from urllib.parse import unquote
from xml.etree import ElementTree as ET

import app as service
import exports
import olefile
import pytest
from ole_package import package_file
from openpyxl import load_workbook
from PIL import Image
from test_app import clients as clients
from test_app import make_project


def unpack_native(content):
    """Independent CFB reader; decode the MS-OLEDS length-delimited payload."""
    with olefile.OleFileIO(io.BytesIO(content), raise_defects=olefile.DEFECT_INCORRECT) as ole:
        # olefile's runtime diagnostics are not declared in its third-party stubs.
        assert not ole.parsing_issues  # pyright: ignore[reportAttributeAccessIssue]
        assert ole.root is not None
        assert ole.root.clsid.upper() == "0003000C-0000-0000-C000-000000000046"
        stream = io.BytesIO(ole.openstream("\x01Ole10Native").read())
        native_size = struct.unpack("<I", stream.read(4))[0]
        assert struct.unpack("<H", stream.read(2))[0] == 2

        def cstring():
            out = bytearray()
            while True:
                char = stream.read(1)
                assert char
                if char == b"\x00":
                    return out.decode("ascii")
                out.extend(char)

        filename = cstring()
        assert cstring() == filename
        stream.read(8)
        assert cstring() == filename
        size = struct.unpack("<I", stream.read(4))[0]
        payload = stream.read(size)
        assert len(payload) == size
        assert stream.tell() == native_size + 4
        return filename, payload


@pytest.mark.parametrize("size", [0, 15, 4096, 8 * 1024 * 1024, 20 * 1024 * 1024])
def test_ole_native_exact_roundtrip_and_large_difat(size):
    original = b"X" * size
    filename, extracted = unpack_native(package_file("original.pdf", original))
    assert filename == "original.pdf"
    assert extracted == original


def upload(client, pid, filename, content):
    response = client.post(f"/api/projects/{pid}/files/file", files={"file": (filename, content)})
    assert response.status_code == 201, response.text
    return response.json()["id"]


def prepare(clients):
    root, student, other = clients
    p, _ = make_project(root)
    image = io.BytesIO()
    Image.new("RGB", (80, 60), "green").save(image, "PNG")
    originals = {}
    for client, attachments in [
        (student, [("图片.png", image.getvalue()), ("同名文档.pdf", b"%PDF-a")]),
        (
            other,
            [
                ("同名文档.pdf", b"%PDF-b"),
                ("报告.docx", b"original-word"),
                ("展示.pptx", b"original-powerpoint"),
            ],
        ),
    ]:
        ids = []
        for name, payload in attachments:
            fid = upload(client, p["id"], name, payload)
            originals[fid] = payload
            ids.append(fid)
        assert (
            client.put(
                f"/api/projects/{p['id']}/submission",
                json={"answers": {"name": "=SUM(1,2)", "file": ids}},
            ).status_code
            == 200
        )
    upload(student, p["id"], "未提交.pdf", b"do-not-export")
    return p, originals


def test_excel_real_images_ole_payloads_and_relationships(clients):
    root, student, _ = clients
    p, originals = prepare(clients)
    assert student.get(f"/api/projects/{p['id']}/export").status_code == 403
    response = root.get(f"/api/projects/{p['id']}/export")
    assert response.status_code == 200, response.text
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        embedded = [n for n in archive.namelist() if n.startswith("xl/embeddings/")]
        assert len(embedded) == 5
        recovered = [unpack_native(archive.read(n))[1] for n in embedded]
        assert sorted(hashlib.sha256(x).hexdigest() for x in recovered) == sorted(
            hashlib.sha256(x).hexdigest() for x in originals.values()
        )
        sheet = ET.fromstring(archive.read("xl/worksheets/sheet2.xml"))
        objects = sheet.findall(f"{{{exports.MAIN}}}oleObjects/{{{exports.MAIN}}}oleObject")
        assert len(objects) == 5
        vml = ET.fromstring(archive.read("xl/drawings/attachments.vml"))
        shapes = {n.get("id") for n in vml.findall(f"{{{exports.VML}}}shape")}
        assert all(f"_x0000_s{n.get('shapeId')}" in shapes for n in objects)
        for name in archive.namelist():
            if name.endswith(".rels"):
                base = posixpath.dirname(posixpath.dirname(name))
                for rel in ET.fromstring(archive.read(name)):
                    assert rel.get("TargetMode") != "External", (name, rel.attrib)
                    target = rel.get("Target", "")
                    resolved = (
                        target.lstrip("/")
                        if target.startswith("/")
                        else posixpath.normpath(posixpath.join(base, target))
                    )
                    assert resolved in archive.namelist(), (name, resolved)
        assert all(
            b"/api/files/" not in archive.read(n) for n in archive.namelist() if n.endswith(".xml")
        )
    book = load_workbook(io.BytesIO(response.content))
    assert book["填写结果"].cell(2, 4).data_type == "s"
    assert book["填写结果"].cell(2, 4).value == "=SUM(1,2)"
    # Runtime-only image list is absent from openpyxl's third-party type stubs.
    assert len(book["附件"]._images) == 1  # pyright: ignore[reportAttributeAccessIssue]
    link = book["填写结果"].cell(2, 5).hyperlink
    assert link is not None and link.location is not None
    assert link.location.startswith("'附件'!")


def test_zip_preserves_originals_and_offline_paths(clients):
    root, student, _ = clients
    p, originals = prepare(clients)
    assert student.get(f"/api/projects/{p['id']}/export?format=zip").status_code == 403
    response = root.get(f"/api/projects/{p['id']}/export?format=zip")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        manifest = json.loads(archive.read("manifest.json"))
        assert len(manifest) == len(originals)
        paths = [item["path"] for item in manifest]
        assert len(set(paths)) == len(paths)
        for item in manifest:
            path = PurePosixPath(item["path"])
            assert not path.is_absolute() and ".." not in path.parts
            assert archive.read(item["path"]) == originals[item["id"]]
            assert hashlib.sha256(archive.read(item["path"])).hexdigest() == item["sha256"]
        book = load_workbook(io.BytesIO(archive.read("workbook.xlsx")))
        assert len(book["附件"]._images) == 1  # pyright: ignore[reportAttributeAccessIssue]
        for row in range(2, len(manifest) + 2):
            link = book["附件"].cell(row, 7).hyperlink
            assert link is not None and link.target is not None
            assert unquote(link.target) in paths


def test_missing_files_and_export_limits_fail_explicitly(clients, monkeypatch):
    root, student, _ = clients
    p, _ = make_project(root)
    fid = upload(student, p["id"], "a.pdf", b"payload")
    student.put(
        f"/api/projects/{p['id']}/submission", json={"answers": {"name": "x", "file": [fid]}}
    )
    with monkeypatch.context() as m:
        m.setattr(exports, "MAX_EXPORT_BYTES", 1)
        assert root.get(f"/api/projects/{p['id']}/export").status_code == 413
    with monkeypatch.context() as m:
        m.setattr(exports, "MAX_EXPORT_FILES", 0)
        assert root.get(f"/api/projects/{p['id']}/export?format=zip").status_code == 413
    (service.DATA / "uploads" / fid).unlink()
    assert root.get(f"/api/projects/{p['id']}/export").status_code == 409
    assert root.get(f"/api/projects/{p['id']}/export?format=zip").status_code == 409


def test_corrupt_images_remain_embedded_and_names_are_safe(clients):
    root, student, _ = clients
    p, _ = make_project(root)
    fid = upload(student, p["id"], "=bad <image>.png", b"not-an-image")
    student.put(
        f"/api/projects/{p['id']}/submission", json={"answers": {"name": "x", "file": [fid]}}
    )
    response = root.get(f"/api/projects/{p['id']}/export")
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert unpack_native(archive.read("xl/embeddings/attachment1.bin"))[1] == b"not-an-image"
    book = load_workbook(io.BytesIO(response.content))
    assert book["附件"].cell(2, 4).data_type == "s"
    assert "无法生成图片预览" in str(book["附件"].cell(2, 6).value)
    assert "/" not in exports.portable_name("../a/b.pdf")
    assert "\\" not in exports.portable_name("..\\a\\b.pdf")


def test_empty_export_still_opens(clients):
    root, _, _ = clients
    p, _ = make_project(root)
    response = root.get(f"/api/projects/{p['id']}/export")
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    assert book["填写结果"].max_row == 1
    assert book["附件"].max_row == 1
