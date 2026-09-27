import io
import re
import zipfile
from urllib.parse import unquote

import app as service
import exports
import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from test_app import clients as clients
from test_app import make_project


def test_class_settings_permissions_and_persistence(clients):
    root, student, _ = clients
    assert root.get("/api/class-settings").json() == {"class_name": "我们的班级"}
    assert student.put("/api/class-settings", json={"class_name": "冒用名称"}).status_code == 403
    response = root.put("/api/class-settings", json={"class_name": "  2026级计算机一班  "})
    assert response.status_code == 200
    assert response.json()["class_name"] == "2026级计算机一班"
    assert student.get("/api/class-settings").json()["class_name"] == "2026级计算机一班"
    service.initialize()
    assert root.get("/api/class-settings").json()["class_name"] == "2026级计算机一班"
    assert root.get("/api/ai/config").json()["model"] == "gpt-4o-mini"
    for invalid in ["", "   ", "a\nname", "\u0000", "x" * 81]:
        assert root.put("/api/class-settings", json={"class_name": invalid}).status_code == 422
    assert root.get("/api/class-settings").json()["class_name"] == "2026级计算机一班"


def test_admin_profile_updates_identity_not_answers_or_privileges(clients):
    root, student, _ = clients
    p, _ = make_project(root)
    url = f"/api/projects/{p['id']}/submission"
    assert root.put(url, json={"answers": {"name": "手动填写的答案"}}).status_code == 200
    student_id = student.get("/api/auth/me").json()["id"]
    assert (
        student.put("/api/admin/profile", json={"name": "不能自行更改审核身份"}).status_code == 403
    )
    response = root.put(
        "/api/admin/profile",
        json={
            "name": "  陈小雨  ",
            "student_no": "20260001",
            "user_id": student_id,
            "role": "student",
        },
    )
    assert response.status_code == 200
    updated = response.json()
    assert updated["name"] == "陈小雨" and updated["student_no"] == "20260001"
    assert updated["role"] == "admin" and updated["username"] == "admin"
    assert "password" not in updated
    assert root.get("/api/auth/me").json()["name"] == "陈小雨"
    assert student.get("/api/auth/me").json()["name"] == "student"
    row = root.get(f"/api/projects/{p['id']}/submissions").json()[0]
    assert row["name"] == "陈小雨" and row["student_no"] == "20260001"
    assert row["answers"]["name"] == "手动填写的答案"
    assert (
        root.put("/api/admin/profile", json={"name": "新姓名"}).json()["student_no"] == "20260001"
    )
    assert (
        root.put("/api/admin/profile", json={"name": "新姓名", "student_no": ""}).json()[
            "student_no"
        ]
        == ""
    )
    for invalid in ["", "  ", "a\nname", "\u0000", "x" * 41]:
        assert root.put("/api/admin/profile", json={"name": invalid}).status_code == 422
    assert root.get("/api/auth/me").json()["name"] == "新姓名"


def test_export_names_class_prefix_and_admin_identity(clients):
    root, _, _ = clients
    p, _ = make_project(root, title="新学期信息登记")
    root.put("/api/class-settings", json={"class_name": "2026级计算机一班"})
    root.put("/api/admin/profile", json={"name": "陈小雨", "student_no": "20260001"})
    root.put(f"/api/projects/{p['id']}/submission", json={"answers": {"name": "答案保留"}})
    expected = "2026级计算机一班_新学期信息登记"
    for extension in ["xlsx", "zip"]:
        response = root.get(f"/api/projects/{p['id']}/export?format={extension}")
        assert response.status_code == 200
        header = response.headers["content-disposition"]
        assert unquote(header.split("filename*=UTF-8''", 1)[1]) == expected + "." + extension
        assert header.isascii() and "\r" not in header and "\n" not in header
        content = response.content
        if extension == "zip":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                assert expected + ".xlsx" in archive.namelist()
                assert expected in archive.read("README.txt").decode()
                content = archive.read(expected + ".xlsx")
        book = load_workbook(io.BytesIO(content))
        sheet = book["填写结果"]
        assert isinstance(sheet, Worksheet)
        assert sheet.cell(2, 1).value == "陈小雨"
        assert sheet.cell(2, 2).value == "20260001"
        info = book["使用说明"]
        assert isinstance(info, Worksheet)
        assert info.cell(2, 2).value == "2026级计算机一班"
    root.put("/api/class-settings", json={"class_name": "新的班级"})
    response = root.get(f"/api/projects/{p['id']}/export")
    assert "新的班级_新学期信息登记.xlsx" in unquote(response.headers["content-disposition"])


def test_export_filename_sanitization_and_length():
    filename = exports.export_filename('../班级\\名称\r\n"', "../项目:<名称>?\\x\u202e")
    assert not re.search(r'[\\/:*?"<>|\x00-\x1f\x7f\u202a-\u202e]', filename)
    assert not filename.startswith(".") and filename.endswith(".xlsx")
    long_name = exports.export_filename("班" * 80, "题" * 150)
    assert len(long_name.encode("utf-8")) <= 240
    assert long_name.encode("utf-8").decode("utf-8") == long_name
    assert exports.export_filename("  ", "...", "zip") == "班级_项目.zip"
    with pytest.raises(ValueError):
        exports.export_filename("班级", "项目", "exe")
